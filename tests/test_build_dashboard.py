"""Task 7 RED: a disposable, local-only XLSX projection of Task 6 evidence.

Run with the bundled Python (openpyxl is required). Tables are located by
semantic headers, never coordinates, colours, column order, or exact widths.
Synthetic summary metrics intentionally disagree with tracker-derived rates.
"""

import copy
import csv
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unittest
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts/build_dashboard.py"
SUMMARY = "reports/analytics-summary.json"
SOURCES = ("tracker/applications.csv", "tracker/events.jsonl", SUMMARY)
OUTPUT = "dashboard/dashboard.xlsx"
HEADER = (
    "canonical_job_id", "jd_snapshot_path", "source", "canonical_url",
    "job_title", "company", "location", "discovery_status",
    "verification_status", "screening_status", "application_stage",
    "action_status", "first_seen_at", "last_verified_at", "applied_at",
    "interview_at", "offer_at", "first_response_at", "next_action",
    "action_due_at", "blocker_category",
)
STATUS_FIELDS = HEADER[7:12]
METRICS = (
    "application_rate", "employer_response_rate",
    "application_to_interview_rate", "interview_to_offer_rate",
    "application_to_offer_rate",
)
LOCAL_ZONE = ZoneInfo("Australia/Melbourne")

# Process-boundary fault injection: production CLI code runs unchanged, using
# this interpreter, with audited private-file reads and network access blocked.
RUNNER = r'''
import os, pathlib, runpy, sys
script, root, fail_replace = sys.argv[1:]
root = pathlib.Path(root).resolve()
allowed = {root / p for p in (
    "tracker/applications.csv", "tracker/events.jsonl",
    "reports/analytics-summary.json")}
violations = []
def reject(message):
    violations.append(message)
    raise PermissionError(message)
def audit(event, args):
    if event.startswith("socket.") or event in ("subprocess.Popen", "os.system"):
        reject("network/process access forbidden")
    if event == "open" and isinstance(args[0], (str, bytes, os.PathLike)):
        path = pathlib.Path(os.fsdecode(args[0])).resolve()
        if path.is_relative_to(root.parent):
            mode, flags = args[1], args[2]
            writing = bool(flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND))
            if writing and not path.is_relative_to(root / "dashboard"):
                reject("write outside dashboard: " + str(path))
            if not writing and path not in allowed:
                reject("read outside three dashboard inputs: " + str(path))
sys.addaudithook(audit)
original_replace = os.replace
def replace(source, destination, *args, **kwargs):
    if fail_replace == "yes" and pathlib.Path(destination).resolve() == root / "dashboard/dashboard.xlsx":
        print("INJECTED_DASHBOARD_REPLACE_FAILURE", file=sys.stderr)
        raise OSError("injected dashboard replacement failure")
    return original_replace(source, destination, *args, **kwargs)
os.replace = replace
sys.path.insert(0, str(pathlib.Path(script).parent))
sys.argv = [script, "--workspace", str(root)]
try:
    runpy.run_path(script, run_name="__main__")
finally:
    if violations:
        print("DASHBOARD_IO_VIOLATION: " + repr(violations), file=sys.stderr)
        raise SystemExit(97)
'''


def normal(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").lower()).strip("_")


def fingerprint(path):
    stat = path.stat()
    return (hashlib.sha256(path.read_bytes()).hexdigest(), stat.st_ino,
            stat.st_mode, stat.st_mtime_ns, stat.st_size)


def snapshot(directory):
    return {str(p.relative_to(directory)): fingerprint(p)
            for p in directory.rglob("*") if p.is_file()}


def text_of(sheet):
    return "\n".join(str(cell.value) for row in sheet for cell in row
                     if cell.value is not None)


def table(sheet, required):
    """Find semantic columns; accept display labels as well as snake_case.

    required maps logical names to accepted, equivalent header labels.
    Repeated tables/sections may follow; callers select records by identifiers.
    """
    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
        labels = {normal(value): column for column, value in enumerate(row)
                  if value is not None}
        columns = {}
        for field, aliases in required.items():
            match = next((labels[normal(name)] for name in (field, *aliases)
                          if normal(name) in labels), None)
            if match is not None:
                columns[field] = match
        if len(columns) == len(required):
            return [dict((key, values[column]) for key, column in columns.items())
                    for values in sheet.iter_rows(min_row=index + 1, values_only=True)
                    if any(value is not None for value in values)]
    raise AssertionError(f"{sheet.title}: missing semantic table columns {list(required)}")


def date_value(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str):
        try:
            parsed = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
            if parsed.tzinfo is not None:
                return parsed.astimezone(LOCAL_ZONE).date().isoformat()
        except ValueError:
            pass
    return str(value or "")[:10]


class DashboardTests(unittest.TestCase):
    maxDiff = None

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="dashboard-contract-")
        self.addCleanup(self.temporary.cleanup)
        self.parent = Path(self.temporary.name)
        self.workspace = self.parent / "local-candidate"
        (self.workspace / "tracker").mkdir(parents=True)
        (self.workspace / "reports").mkdir()
        self.today = dt.datetime.now(LOCAL_ZONE).date()
        self.rows = []
        self.events = []
        cases = (
            ("due-today", "submitted", "follow_up", 0, "Check receipt", ""),
            ("overdue", "interview", "follow_up", -2, "Confirm interview", ""),
            ("approval", "ready", "candidate_review", None, "Approve final documents", ""),
            ("blocked", "preparing", "blocked", None, "Resolve missing evidence", "missing_evidence"),
            ("undated", "submitted", "follow_up", None, "Ask for an update", ""),
            ("future", "shortlisted", "follow_up", 4, "Check upcoming vacancy", ""),
            ("offer", "offer", "none", None, "", ""),
        )
        for index, (identifier, stage, action, offset, next_action, blocker) in enumerate(cases):
            row = dict.fromkeys(HEADER, "")
            row.update(canonical_job_id=identifier,
                       jd_snapshot_path=f"jobs/snapshots/{identifier}.md",
                       source="board_a" if index % 2 == 0 else "board_b",
                       canonical_url=f"https://jobs.example.test/{identifier}",
                       job_title="Synthetic Analyst", company="Fictional Employer",
                       location="Melbourne VIC", discovery_status="discovered",
                       verification_status="jd_verified", screening_status="eligible",
                       application_stage=stage, action_status=action,
                       first_seen_at="2026-07-01T10:00:00+10:00",
                       last_verified_at="2026-07-02T10:00:00+10:00",
                       next_action=next_action, blocker_category=blocker)
            if offset is not None:
                # Early local morning is the previous UTC calendar date. This
                # catches comparing raw UTC date/string slices to local today.
                due = dt.datetime.combine(self.today + dt.timedelta(days=offset),
                                          dt.time(0, 30), LOCAL_ZONE)
                row["action_due_at"] = due.astimezone(dt.timezone.utc).isoformat()
            if stage in ("submitted", "interview", "offer"):
                row["applied_at"] = "2026-07-06T10:00:00+10:00"
            if stage in ("interview", "offer"):
                row["interview_at"] = "2026-07-14T10:00:00+10:00"
            if stage == "offer":
                row["offer_at"] = "2026-07-22T10:00:00+10:00"
            self.rows.append(row)
            self.events.append(self.event(row, f"{identifier}-recorded", "application_recorded"))
        target = self.rows[0]
        self.events.extend([
            self.event(target, "docs-v1", "documents_updated", versions={"resume": "resume-v1"}),
            self.event(target, "docs-v2", "documents_updated", versions={"resume": "resume-v2", "cover_letter": "cover-v3"}),
            self.event(target, "after-docs-note", "note_recorded", reason="Receipt needs checking"),
            self.event(self.rows[4], "undated-reason-event", "note_recorded",
                       reason="Recruiter requested contact after internal review; no date agreed"),
        ])
        self.summary = {
            "as_of": "2026-09-15T00:00:00Z",
            "sources": list(SOURCES[:2]),
            "overall": self.metrics(1, 20),
            "segments": {"source": {
                "board_a": {"metrics": self.metrics(2, 25)},
                "board_b": {"metrics": self.metrics(1, 8)},
            }},
            "ranked_conclusions": [],
        }
        self.write_sources()
        # Round-trip fixture validation without importing production validators
        # (update_tracker imports PyYAML, which bundled Python need not have).
        with (self.workspace / SOURCES[0]).open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            self.assertEqual(tuple(reader.fieldnames), HEADER)
            self.assertEqual(len(HEADER), 21)
            self.assertEqual(list(reader), self.rows)
        emitted = [json.loads(line) for line in
                   (self.workspace / SOURCES[1]).read_text(encoding="utf-8").splitlines()]
        self.assertEqual(emitted, self.events)
        event_fields = {
            "event_id", "occurred_at", "canonical_job_id", "event_type",
            "from_status", "to_status", "jd_snapshot_path", "document_versions",
            "source", "reason", *HEADER[17:],
        }
        for event in emitted:
            self.assertEqual(set(event), event_fields)
            self.assertEqual(set(event["to_status"]), set(STATUS_FIELDS))
            self.assertIsNotNone(dt.datetime.fromisoformat(event["occurred_at"]).tzinfo)
        for row in self.rows:
            for field in ("first_seen_at", "last_verified_at", "applied_at", "interview_at",
                          "offer_at", "first_response_at", "action_due_at"):
                if row[field]:
                    self.assertIsNotNone(dt.datetime.fromisoformat(row[field]).tzinfo)

    @staticmethod
    def metrics(start, denominator):
        return {name: {
            "numerator": start + index, "denominator": denominator,
            "maturity_window": {"days": (0, 30, 30, 30, 60)[index],
                                "terminal_outcome_is_mature": True},
            "rate": (start + index) / denominator,
            "data_status": "sufficient_data" if denominator >= 10 else "insufficient_data",
        } for index, name in enumerate(METRICS)}

    @staticmethod
    def event(row, identifier, kind, *, versions=None, reason=""):
        status = {field: row[field] for field in STATUS_FIELDS}
        return {
            "event_id": identifier, "occurred_at": "2026-08-01T10:00:00+10:00",
            "canonical_job_id": row["canonical_job_id"], "event_type": kind,
            "from_status": None if kind == "application_recorded" else dict(status),
            "to_status": status, "jd_snapshot_path": row["jd_snapshot_path"],
            "document_versions": versions or {}, "source": row["source"], "reason": reason,
            **{field: row[field] for field in HEADER[17:]},
        }

    def write_sources(self):
        with (self.workspace / SOURCES[0]).open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=HEADER)
            writer.writeheader()
            writer.writerows(self.rows)
        (self.workspace / SOURCES[1]).write_text(
            "".join(json.dumps(event) + "\n" for event in self.events), encoding="utf-8")
        (self.workspace / SUMMARY).write_text(json.dumps(self.summary), encoding="utf-8")

    def run_builder(self, *, fail_replace=False):
        before = snapshot(self.parent)
        self.assertTrue(SCRIPT.is_file(), f"Task 7 dashboard script is absent: {SCRIPT}")
        result = subprocess.run(
            [sys.executable, "-B", "-c", RUNNER, str(SCRIPT), str(self.workspace),
             "yes" if fail_replace else "no"],
            cwd=ROOT, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1", "TZ": "Australia/Melbourne"},
            text=True, capture_output=True, timeout=30, check=False)
        self.assertNotIn("DASHBOARD_IO_VIOLATION", result.stderr, result.stderr)
        after = snapshot(self.parent)
        output = "local-candidate/" + OUTPUT
        self.assertEqual({p: v for p, v in before.items() if p != output},
                         {p: v for p, v in after.items() if p != output},
                         "Only dashboard/dashboard.xlsx may change; no temporary files may remain")
        return result

    def build(self):
        result = self.run_builder()
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        workbook = load_workbook(self.workspace / OUTPUT, data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def test_exact_sheets_order_and_only_one_generated_file(self):
        self.assertEqual(self.build().sheetnames,
                         ["Overview", "Applications", "Conversion", "Follow-ups"])

    def test_reads_only_three_local_sources_and_never_network_or_other_candidate(self):
        (self.workspace / "profile").mkdir()
        (self.workspace / "profile/candidate.yaml").write_text("malformed: [", encoding="utf-8")
        (self.workspace / "profile/analytics-consent.yaml").write_text("malformed: [", encoding="utf-8")
        neighbor = self.parent / "other-candidate/tracker"
        neighbor.mkdir(parents=True)
        (neighbor / "applications.csv").write_text("MUST_NOT_READ_NEIGHBOR", encoding="utf-8")
        workbook = self.build()
        self.assertNotIn("MUST_NOT_READ_NEIGHBOR", "\n".join(text_of(s) for s in workbook))

    def test_overview_validated_counts_have_source_references(self):
        sheet = self.build()["Overview"]
        expected = {
            "total": (("total_applications", "tracked_jobs", "total_jobs", "applications"), 7),
            "verified": (("verified_jobs", "jd_verified", "verified"), 7),
            "approval": (("approval_needs", "needs_approval", "candidate_review", "approvals"), 1),
            "blocked": (("blocked_actions", "blocked"), 1),
            "overdue": (("overdue_actions", "overdue"), 1),
            "today": (("due_today", "actions_due_today"), 1),
        }
        rows = list(sheet.iter_rows(values_only=True))
        for label, (aliases, count) in expected.items():
            with self.subTest(count=label):
                matches = [r for r in rows if any(normal(c) in aliases for c in r)]
                self.assertTrue(any(count in r and SOURCES[0] in " ".join(map(str, r))
                                    for r in matches), f"Missing source-traceable {label} count {count}")

    def test_overview_priorities_approvals_blockers_and_real_due_dates(self):
        sheet = self.build()["Overview"]
        records = table(sheet, {
            "canonical_job_id": ("job_id",), "action_due_at": ("due_date", "due_at"),
            "next_action": (), "category": ("priority", "action_category", "reason"),
        })
        by_id = {r["canonical_job_id"]: r for r in records}
        for identifier, category in (("due-today", "today"), ("overdue", "overdue"),
                                     ("approval", "approval|review"), ("blocked", "blocked")):
            self.assertIn(identifier, by_id)
            self.assertRegex(str(by_id[identifier]["category"]).lower(), category)
        for identifier, offset in (("due-today", 0), ("overdue", -2)):
            self.assertEqual(date_value(by_id[identifier]["action_due_at"]),
                             (self.today + dt.timedelta(days=offset)).isoformat())
        for identifier in ("undated", "approval", "blocked"):
            if identifier in by_id:
                self.assertIn(by_id[identifier]["action_due_at"], (None, "", "n/a", "No due date"))
        if "future" in by_id:
            self.assertNotRegex(str(by_id["future"]["category"]).lower(), "overdue|due_today|due today")

    def test_overview_conversion_rates_are_copied_from_summary(self):
        sheet = self.build()["Overview"]
        records = table(sheet, {"metric": (), "rate": (), "source_path": ("source", "evidence_source")})
        for name, metric in self.summary["overall"].items():
            matches = [r for r in records if normal(r["metric"]) == name]
            self.assertTrue(any(r["rate"] == metric["rate"] and r["source_path"] == SUMMARY
                                for r in matches), f"Overview must use summary {name}")

    def test_applications_exposes_identity_five_statuses_documents_and_sources(self):
        sheet = self.build()["Applications"]
        fields = {field: () for field in ("canonical_job_id", *STATUS_FIELDS,
                                          "jd_snapshot_path", "source", "canonical_url")}
        records = table(sheet, fields)
        by_id = {r["canonical_job_id"]: r for r in records}
        self.assertEqual(set(by_id), {r["canonical_job_id"] for r in self.rows})
        for row in self.rows:
            for field in fields:
                self.assertEqual(by_id[row["canonical_job_id"]][field], row[field])
        job_rows = [r for r in sheet.iter_rows(values_only=True) if "due-today" in r]
        job_text = " ".join(str(c) for r in job_rows for c in r)
        self.assertIn("resume-v2", job_text)
        self.assertIn("cover-v3", job_text)
        self.assertIn(SOURCES[1], text_of(sheet))

    def test_conversion_copies_overall_and_source_metrics_with_evidence(self):
        sheet = self.build()["Conversion"]
        records = table(sheet, {
            "metric": (), "numerator": (), "denominator": (),
            "maturity_window_days": ("maturity_days", "maturity_window"),
            "rate": (), "data_status": (), "source_path": ("evidence_source", "analytics_source"),
        })
        # Accept either segment columns or labeled sections, but require the
        # correct metric values to belong to the correct segment.
        contextual_values = {}
        context = None
        for values in sheet.iter_rows(values_only=True):
            labels = {normal(value) for value in values if value is not None}
            for segment in ("overall", "board_a", "board_b"):
                if labels & {segment, "source_" + segment, "segment_" + segment}:
                    context = segment
            for metric_name in METRICS:
                if metric_name in labels:
                    contextual_values.setdefault((context, metric_name), []).append(values)
        for segment, metrics in [("overall", self.summary["overall"]), *[
            (source, payload["metrics"]) for source, payload in self.summary["segments"]["source"].items()
        ]]:
            self.assertIn(segment, text_of(sheet).lower())
            for name, metric in metrics.items():
                matches = [r for r in records if normal(r["metric"]) == name
                           and r["numerator"] == metric["numerator"]
                           and r["denominator"] == metric["denominator"]]
                self.assertEqual(len(matches), 1, (segment, name, matches))
                record = matches[0]
                self.assertTrue(any(metric["numerator"] in values
                                    and metric["denominator"] in values
                                    and metric["rate"] in values
                                    for values in contextual_values.get((segment, name), [])),
                                f"Metric is not associated with its source segment: {segment}/{name}")
                self.assertEqual(record["maturity_window_days"], metric["maturity_window"]["days"])
                self.assertAlmostEqual(record["rate"], metric["rate"])
                self.assertEqual(record["data_status"], metric["data_status"])
                self.assertEqual(record["source_path"], SUMMARY)

    def test_conversion_weekly_counts_use_actual_funnel_timestamps(self):
        records = table(self.build()["Conversion"], {
            "week_start": ("week", "week_commencing"),
            "applications": ("submitted", "applications_count", "application_count"),
            "interviews": ("interviews_count", "interview_count"),
            "offers": ("offers_count", "offer_count"),
        })
        weeks = {date_value(r["week_start"]): r for r in records}
        for week, counts in {"2026-07-06": (4, 0, 0), "2026-07-13": (0, 2, 0),
                             "2026-07-20": (0, 0, 1)}.items():
            self.assertIn(week, weeks)
            self.assertEqual(tuple(weeks[week][key] for key in ("applications", "interviews", "offers")), counts)

    def test_followups_link_current_state_and_event_reason_without_inventing_dates(self):
        sheet = self.build()["Follow-ups"]
        records = table(sheet, {
            "canonical_job_id": (), "action_due_at": ("due_date", "due_at"),
            "application_stage": ("stage",), "action_status": (), "next_action": (),
            "blocker_category": (), "reason": ("event_reason",),
            "event_ids": ("supporting_event_ids", "event_id"),
        })
        by_id = {r["canonical_job_id"]: r for r in records}
        for row in self.rows:
            if row["action_status"] == "none":
                continue
            identifier = row["canonical_job_id"]
            self.assertIn(identifier, by_id)
            actual = by_id[identifier]
            for key in ("application_stage", "action_status", "next_action", "blocker_category"):
                self.assertEqual(actual[key] or "", row[key])
            if row["action_due_at"]:
                expected = dt.datetime.fromisoformat(row["action_due_at"]).astimezone(LOCAL_ZONE).date()
                self.assertEqual(date_value(actual["action_due_at"]), expected.isoformat())
            else:
                self.assertIn(actual["action_due_at"], (None, "", "n/a", "No due date"))
            supporting = [e["event_id"] for e in self.events if e["canonical_job_id"] == identifier]
            self.assertTrue(any(e in str(actual["event_ids"]) for e in supporting))
        self.assertIn("undated-reason-event", str(by_id["undated"]["event_ids"]))
        self.assertIn(self.events[-1]["reason"], str(by_id["undated"]["reason"]))
        self.assertIn(SOURCES[1], text_of(sheet))

    def test_append_order_controls_latest_evidence_with_backdated_events(self):
        target = self.rows[0]
        latest = self.event(
            target, "backdated-docs", "documents_updated",
            versions={"resume": "resume-v4"}, reason="Use the corrected final documents",
        )
        latest["occurred_at"] = "2026-07-31T10:00:00+10:00"
        blank_note = self.event(target, "backdated-blank-note", "note_recorded")
        blank_note["occurred_at"] = "2026-07-30T10:00:00+10:00"
        self.events.extend([latest, blank_note])
        self.write_sources()
        workbook = self.build()
        application = next(record for record in table(workbook["Applications"], {
            "canonical_job_id": (), "document_versions": (), "event_ids": (),
        }) if record["canonical_job_id"] == target["canonical_job_id"])
        followup = next(record for record in table(workbook["Follow-ups"], {
            "canonical_job_id": (), "reason": ("event_reason",), "event_ids": (),
        }) if record["canonical_job_id"] == target["canonical_job_id"])
        with self.subTest(evidence="latest document versions"):
            self.assertEqual(json.loads(application["document_versions"]),
                             {"resume": "resume-v4", "cover_letter": "cover-v3"})
        with self.subTest(evidence="latest nonempty reason"):
            self.assertEqual(followup["reason"], latest["reason"])
        supporting = [event["event_id"] for event in self.events
                      if event["canonical_job_id"] == target["canonical_job_id"]]
        for sheet, record in (("Applications", application), ("Follow-ups", followup)):
            with self.subTest(evidence="supporting event append order", sheet=sheet):
                identifiers = str(record["event_ids"])
                for identifier in supporting:
                    self.assertIn(identifier, identifiers)
                positions = [identifiers.index(identifier) for identifier in supporting]
                self.assertEqual(positions, sorted(positions))

    def precreate_workbook(self):
        destination = self.workspace / OUTPUT
        destination.parent.mkdir(exist_ok=True)
        workbook = Workbook()
        workbook.active.title = "Tampered"
        workbook.active["A1"] = "WORKBOOK_EDITS_MUST_NOT_REACH_SOURCES"
        workbook.save(destination)
        workbook.close()
        return destination

    def test_regenerates_tampered_workbook_from_changed_sources(self):
        self.precreate_workbook()
        first = self.build()
        first["Applications"]["A2"] = "WORKBOOK_EDITS_MUST_NOT_REACH_SOURCES"
        first.save(self.workspace / OUTPUT)
        self.rows[0]["job_title"] = "Updated synthetic title from tracker"
        self.events.append(self.event(self.rows[0], "docs-v4", "documents_updated",
                                      versions={"resume": "resume-v4"}))
        self.summary["overall"]["application_rate"].update(numerator=9, rate=9 / 20)
        self.write_sources()
        workbook = self.build()
        content = "\n".join(text_of(s) for s in workbook)
        self.assertNotIn("WORKBOOK_EDITS_MUST_NOT_REACH_SOURCES", content)
        self.assertIn("Updated synthetic title from tracker", content)
        self.assertIn("resume-v4", text_of(workbook["Applications"]))
        rates = table(workbook["Overview"], {"metric": (), "rate": ()})
        self.assertTrue(any(normal(r["metric"]) == "application_rate" and r["rate"] == 9 / 20 for r in rates))

    def assert_rejected_without_workbook(self):
        result = self.run_builder()
        self.assertNotEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertFalse((self.workspace / OUTPUT).exists())

    def test_missing_summary_fails_without_creating_workbook(self):
        (self.workspace / SUMMARY).unlink()
        self.assert_rejected_without_workbook()

    def test_malformed_summary_fails_without_creating_workbook(self):
        (self.workspace / SUMMARY).write_text('{"overall":', encoding="utf-8")
        self.assert_rejected_without_workbook()

    def test_invalid_summary_contract_fails_without_creating_workbook(self):
        mutations = {
            "missing_overall": lambda s: s.pop("overall"),
            "missing_sources": lambda s: s.pop("sources"),
            "missing_source_segments": lambda s: s["segments"].pop("source"),
            "missing_conclusions": lambda s: s.pop("ranked_conclusions"),
            "invalid_rate": lambda s: s["overall"]["application_rate"].update(rate="not-a-rate"),
            "negative_denominator": lambda s: s["overall"]["application_rate"].update(denominator=-1),
            "invalid_maturity": lambda s: s["overall"]["application_rate"].update(maturity_window={"days": -1}),
            "invalid_segment": lambda s: s["segments"]["source"]["board_a"]["metrics"].pop("application_rate"),
        }
        for name, mutate in mutations.items():
            with self.subTest(contract=name):
                summary = copy.deepcopy(self.summary)
                mutate(summary)
                (self.workspace / SUMMARY).write_text(json.dumps(summary), encoding="utf-8")
                self.assert_rejected_without_workbook()

    def summary_maturity_invalid_cases(self):
        # Exact JSON distinguishes booleans/numbers from strings and absent keys.
        windows = (
            ("application_days_1", "application_rate",
             '{"days": 1, "terminal_outcome_is_mature": true}'),
            ("application_days_30", "application_rate",
             '{"days": 30, "terminal_outcome_is_mature": true}'),
            ("application_days_string_zero", "application_rate",
             '{"days": "0", "terminal_outcome_is_mature": true}'),
            ("application_days_missing", "application_rate",
             '{"terminal_outcome_is_mature": true}'),
            ("terminal_false", "employer_response_rate",
             '{"days": 30, "terminal_outcome_is_mature": false}'),
            ("terminal_string_true", "employer_response_rate",
             '{"days": 30, "terminal_outcome_is_mature": "true"}'),
            ("terminal_missing", "employer_response_rate", '{"days": 30}'),
        )
        for scope in ("overall", "source_board_a"):
            for name, metric, window_json in windows:
                summary = copy.deepcopy(self.summary)
                metrics = (summary["overall"] if scope == "overall" else
                           summary["segments"]["source"]["board_a"]["metrics"])
                metrics[metric]["maturity_window"] = json.loads(window_json)
                yield scope, name, summary

    def assert_summary_maturity_rejected(self, summary, *, preserve_existing):
        destination = self.workspace / OUTPUT
        # Reset only the disposable output so one RED subcase cannot contaminate another.
        destination.unlink(missing_ok=True)
        before = None
        if preserve_existing:
            before = self.precreate_workbook().read_bytes()
        (self.workspace / SUMMARY).write_text(json.dumps(summary), encoding="utf-8")
        result = self.run_builder()
        observed = {
            "rejected": result.returncode != 0,
            "temporary_files": list(destination.parent.glob(".dashboard-*")),
        }
        expected = {"rejected": True, "temporary_files": []}
        if preserve_existing:
            observed["original_bytes_preserved"] = (
                destination.exists() and destination.read_bytes() == before)
            observed["old_workbook_marker"] = None
            if destination.exists():
                existing = load_workbook(destination)
                try:
                    if existing.sheetnames == ["Tampered"]:
                        observed["old_workbook_marker"] = existing["Tampered"]["A1"].value
                finally:
                    existing.close()
            expected.update(original_bytes_preserved=True,
                            old_workbook_marker="WORKBOOK_EDITS_MUST_NOT_REACH_SOURCES")
        else:
            observed["new_workbook_created"] = destination.exists()
            expected["new_workbook_created"] = False
        self.assertEqual(observed, expected, result.stdout + result.stderr)

    def test_summary_maturity_invalid_without_existing_workbook(self):
        for scope, case, summary in self.summary_maturity_invalid_cases():
            with self.subTest(scope=scope, contract=case):
                self.assert_summary_maturity_rejected(summary, preserve_existing=False)

    def test_summary_maturity_invalid_preserves_existing_workbook(self):
        for scope, case, summary in self.summary_maturity_invalid_cases():
            with self.subTest(scope=scope, contract=case):
                self.assert_summary_maturity_rejected(summary, preserve_existing=True)

    def test_summary_maturity_legal_overall_and_segments_accept(self):
        for metrics in (self.summary["overall"], *[
            segment["metrics"] for segment in self.summary["segments"]["source"].values()
        ]):
            metrics["application_rate"]["maturity_window"] = json.loads(
                '{"days": 0, "terminal_outcome_is_mature": true}')
            for metric in metrics.values():
                metric["maturity_window"]["terminal_outcome_is_mature"] = json.loads('true')
        self.write_sources()
        workbook = self.build()
        self.assertEqual(workbook.sheetnames,
                         ["Overview", "Applications", "Conversion", "Follow-ups"])
        self.assertEqual(list((self.workspace / "dashboard").glob(".dashboard-*")), [])

    def test_atomic_replace_failure_preserves_existing_workbook_and_cleans_temporary(self):
        destination = self.precreate_workbook()
        before = destination.read_bytes()
        result = self.run_builder(fail_replace=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("INJECTED_DASHBOARD_REPLACE_FAILURE", result.stderr)
        self.assertEqual(destination.read_bytes(), before)
        self.assertEqual(list(destination.parent.iterdir()), [destination])
        with destination.open("rb") as handle:
            existing = load_workbook(handle)
            self.assertEqual(existing.sheetnames, ["Tampered"])
            existing.close()

    def test_workbook_has_usable_tables_styled_headers_and_no_formula_errors(self):
        workbook = self.build()
        for sheet in workbook:
            with self.subTest(sheet=sheet.title):
                populated = [c for row in sheet for c in row if c.value is not None]
                self.assertGreater(len(populated), 1)
                self.assertTrue(sheet.freeze_panes or sheet.tables,
                                "Keep table context visible while scrolling")
                if sheet.title != "Overview":
                    self.assertTrue(sheet.auto_filter.ref or sheet.tables,
                                    "Data sheets need filtering/table affordances")
                self.assertTrue(any(c.font.bold or c.fill.fill_type for c in populated),
                                "Headers should be visually distinguished")
                self.assertTrue(any((d.width or 0) >= 12 for d in sheet.column_dimensions.values()),
                                "Use readable column widths")
                self.assertFalse(any(c.data_type == "e" for c in populated))
                for cell in populated:
                    if cell.data_type == "f":
                        self.assertNotRegex(cell.value, r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A|NUM!|NULL!)")
        cached = load_workbook(self.workspace / OUTPUT, data_only=True)
        try:
            for sheet in cached:
                self.assertFalse(any(c.data_type == "e" for row in sheet for c in row))
        finally:
            cached.close()


if __name__ == "__main__":
    unittest.main()
