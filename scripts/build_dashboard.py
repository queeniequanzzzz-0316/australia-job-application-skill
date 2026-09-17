"""Rebuild a disposable dashboard from exactly three local evidence files.

The tracker supplies current state and weekly counts; the analytics summary
supplies every conversion metric. No workbook, profile, snapshot, document or
neighbouring workspace is read. Workbook text is always literal, never a formula.
"""

import argparse
from collections import defaultdict
import csv
import datetime as dt
import json
import math
import os
from pathlib import Path
import re
import stat
import sys
import tempfile
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCES = ("tracker/applications.csv", "tracker/events.jsonl", "reports/analytics-summary.json")
OUTPUT = "dashboard/dashboard.xlsx"
HEADER = (
    "canonical_job_id", "jd_snapshot_path", "source", "canonical_url",
    "job_title", "company", "location", "discovery_status",
    "verification_status", "screening_status", "application_stage",
    "action_status", "first_seen_at", "last_verified_at", "applied_at",
    "interview_at", "offer_at", "first_response_at", "next_action",
    "action_due_at", "blocker_category",
)
STATUS_DIMENSIONS = {
    "discovery_status": ("discovered", "duplicate_of_known", "expired", "delisted"),
    "verification_status": ("unverified", "jd_retrieved", "jd_verified", "verification_failed"),
    "screening_status": ("new", "verified", "eligible", "excluded", "duplicate", "closed"),
    "application_stage": ("not_started", "shortlisted", "preparing", "ready", "submitted",
                          "interview", "offer", "accepted", "rejected", "withdrawn"),
    "action_status": ("none", "candidate_review", "document_edit", "application_form",
                      "follow_up", "blocked"),
}
TIMESTAMPS = (*HEADER[12:18], "action_due_at")
EVENT_TYPES = ("application_recorded", "employer_response_recorded", "status_changed",
               "snapshot_linked", "documents_updated", "note_recorded")
EVENT_FIELDS = {"event_id", "occurred_at", "canonical_job_id", "event_type", "from_status",
                "to_status", "jd_snapshot_path", "document_versions", "source", "reason",
                *HEADER[17:]}
METRICS = ("application_rate", "employer_response_rate", "application_to_interview_rate",
           "interview_to_offer_rate", "application_to_offer_rate")
METRIC_HEADERS = ("metric", "numerator", "denominator", "maturity_window_days",
                  "terminal_outcome_is_mature", "rate", "data_status", "source_path")
LOCAL_ZONE = ZoneInfo("Australia/Melbourne")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def instant(value):
    require(isinstance(value, str) and re.fullmatch(
        r"\d{4}-\d{2}-\d{2}[Tt]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:[Zz]|[+-]\d{2}:\d{2})", value),
        "Expected a timezone-aware RFC 3339 timestamp.")
    if value[-1:] not in ("Z", "z"):
        require(int(value[-5:-3]) <= 23 and int(value[-2:]) <= 59,
                "Invalid timestamp offset.")
    parsed = dt.datetime.fromisoformat(value.upper().replace("Z", "+00:00"))
    return parsed.astimezone(dt.timezone.utc)


def local_date(value):
    return instant(value).astimezone(LOCAL_ZONE).date() if value else None


def contained(root, path):
    relative = path.relative_to(root)
    require(".." not in relative.parts, "Path escapes workspace.")
    current = root
    require(not current.is_symlink(), "Symlink workspace is forbidden.")
    for part in relative.parts:
        current /= part
        require(not current.is_symlink(), "Symlink input or output is forbidden.")
    require(path.resolve().is_relative_to(root), "Path escapes workspace.")


def workspace_paths(path):
    path = path.expanduser()
    require(".." not in path.parts, "Workspace path must not contain parent traversal.")
    absolute = Path(os.path.abspath(path))
    require(all(not ancestor.is_symlink() for ancestor in (absolute, *absolute.parents)),
            "Symlink workspace or ancestor is forbidden.")
    root = absolute.resolve(strict=True)
    require(root.is_dir(), "Workspace must be a directory.")
    paths = [root / name for name in (*SOURCES, OUTPUT)]
    for item in paths:
        contained(root, item)
    for item in paths[:3]:
        require(item.is_file(), f"Missing or invalid input: {item.name}.")
    require(not paths[-1].exists() or paths[-1].is_file(), "Invalid dashboard destination.")
    return root, paths


def open_input(root, path):
    contained(root, path)
    descriptor = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
    try:
        require(stat.S_ISREG(os.fstat(descriptor).st_mode), "Input must be a regular file.")
        return os.fdopen(descriptor, "r", encoding="utf-8", newline="")
    except BaseException:
        os.close(descriptor)
        raise


def validate_status(status):
    require(isinstance(status, dict) and set(status) == set(STATUS_DIMENSIONS),
            "Invalid Task 6 status shape.")
    for field, vocabulary in STATUS_DIMENSIONS.items():
        require(status[field] in vocabulary, f"Invalid {field}.")


def read_rows(root, path):
    jobs = {}
    with open_input(root, path) as handle:
        reader = csv.reader(handle, strict=True)
        require(tuple(next(reader, ())) == HEADER, "Expected exact 21-column Task 6 CSV header.")
        for cells in reader:
            require(len(cells) == len(HEADER), "Invalid application row width.")
            row = dict(zip(HEADER, cells))
            identifier = row["canonical_job_id"]
            require(bool(identifier.strip()), "Missing canonical_job_id.")
            validate_status({field: row[field] for field in STATUS_DIMENSIONS})
            for field in TIMESTAMPS:
                if row[field]:
                    instant(row[field])
            funnel = [instant(row[field]) for field in HEADER[14:17] if row[field]]
            require(funnel == sorted(funnel), "Invalid funnel timestamp order.")
            require(identifier not in jobs or jobs[identifier] == row,
                    "Conflicting rows for canonical_job_id.")
            jobs[identifier] = row
    return jobs


def read_events(root, path, jobs):
    events = []
    with open_input(root, path) as handle:
        for line in handle:
            if not line.strip():
                continue
            event = json.loads(line)
            require(isinstance(event, dict) and set(event) == EVENT_FIELDS,
                    "Invalid Task 6 event shape.")
            for field in ("event_id", "canonical_job_id", "source", "reason", *HEADER[17:]):
                require(isinstance(event[field], str), f"Invalid event {field}.")
            require(bool(event["event_id"].strip()), "Missing event_id.")
            require(event["canonical_job_id"] in jobs, "Event has no application row.")
            require(event["event_type"] in EVENT_TYPES, "Invalid event_type.")
            instant(event["occurred_at"])
            for field in ("first_response_at", "action_due_at"):
                if event[field]:
                    instant(event[field])
            validate_status(event["to_status"])
            if event["from_status"] is not None:
                validate_status(event["from_status"])
            require(event["jd_snapshot_path"] is None or isinstance(event["jd_snapshot_path"], str),
                    "Invalid event snapshot reference.")
            require(isinstance(event["document_versions"], dict) and all(
                isinstance(value, str) for value in event["document_versions"].values()),
                "Invalid document_versions.")
            events.append(event)
    # Append order controls current evidence, including backdated corrections.
    return events


def validate_metrics(metrics):
    require(isinstance(metrics, dict) and set(metrics) == set(METRICS),
            "Expected all five summary metrics.")
    for name, metric in metrics.items():
        require(isinstance(metric, dict) and set(metric) == {
            "numerator", "denominator", "maturity_window", "rate", "data_status"},
            "Invalid summary metric shape.")
        numerator, denominator = metric["numerator"], metric["denominator"]
        require(type(numerator) is int and type(denominator) is int
                and 0 <= numerator <= denominator, "Invalid metric counts.")
        window = metric["maturity_window"]
        require(isinstance(window, dict) and set(window) == {"days", "terminal_outcome_is_mature"}
                and type(window["days"]) is int and window["days"] >= 0
                and type(window["terminal_outcome_is_mature"]) is bool
                and window["terminal_outcome_is_mature"] is True,
                "Invalid metric maturity window.")
        require(name != "application_rate" or window["days"] == 0,
                "Application rate maturity days must be exactly integer zero.")
        rate = metric["rate"]
        # Validate the supplied rate, but always export its original numeric value.
        require((denominator == 0 and rate is None) or (
            denominator > 0 and type(rate) in (int, float) and math.isfinite(rate)
            and 0 <= rate <= 1 and math.isclose(rate, numerator / denominator,
                                               rel_tol=1e-12, abs_tol=1e-15)),
            "Invalid summary rate.")
        require(metric["data_status"] == (
            "insufficient_data" if denominator < 10 else "sufficient_data"),
            "Invalid metric data_status.")


def read_summary(root, path):
    with open_input(root, path) as handle:
        summary = json.load(handle)
    require(isinstance(summary, dict) and set(summary) == {
        "as_of", "sources", "overall", "segments", "ranked_conclusions"},
        "Invalid analytics summary shape.")
    instant(summary["as_of"])
    require(summary["sources"] == list(SOURCES[:2]), "Invalid summary sources.")
    validate_metrics(summary["overall"])
    segments = summary["segments"]
    require(isinstance(segments, dict) and set(segments) == {"source"}
            and isinstance(segments["source"], dict), "Invalid source segments.")
    for segment in segments["source"].values():
        require(isinstance(segment, dict) and set(segment) == {"metrics"},
                "Invalid source segment.")
        validate_metrics(segment["metrics"])
    require(isinstance(summary["ranked_conclusions"], list), "Invalid ranked_conclusions.")
    return summary


def append_row(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        # Prevent input text such as '=...' or '#REF!' becoming executable/error cells.
        if isinstance(cell.value, str):
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(cell.value, dt.date):
            cell.number_format = "yyyy-mm-dd"


def section(sheet, title, headers, rows, *, filtering=False):
    if sheet.max_row > 1 or sheet.cell(1, 1).value is not None:
        append_row(sheet, [])
    append_row(sheet, [title])
    sheet.cell(sheet.max_row, 1).font = Font(bold=True, size=14, color="17365D")
    append_row(sheet, headers)
    header_row = sheet.max_row
    for cell in sheet[header_row][:len(headers)]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A73")
    for values in rows:
        append_row(sheet, values)
        for index, header in enumerate(headers, 1):
            if header == "rate":
                sheet.cell(sheet.max_row, index).number_format = "0.0%"
    if filtering:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{sheet.max_row}"
    if not sheet.freeze_panes:
        sheet.freeze_panes = f"B{header_row + 1}"
    for index, header in enumerate(headers, 1):
        width = 40 if header in ("next_action", "reason", "document_versions", "event_ids") else 24
        if header.endswith("path") or header.endswith("reference") or header == "canonical_url":
            width = 38
        dimension = sheet.column_dimensions[get_column_letter(index)]
        dimension.width = max(dimension.width or 0, width)
    sheet.sheet_view.showGridLines = False


def metric_rows(metrics):
    for name in METRICS:
        value = metrics[name]
        yield (name, value["numerator"], value["denominator"],
               value["maturity_window"]["days"],
               value["maturity_window"]["terminal_outcome_is_mature"],
               value["rate"], value["data_status"], SOURCES[2])


def build_workbook(jobs, events, summary):
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    applications = workbook.create_sheet("Applications")
    conversion = workbook.create_sheet("Conversion")
    followups = workbook.create_sheet("Follow-ups")
    today = dt.datetime.now(LOCAL_ZONE).date()
    rows = list(jobs.values())
    history = defaultdict(lambda: {"versions": {}, "reason": "", "ids": []})
    for event in events:
        evidence = history[event["canonical_job_id"]]
        evidence["versions"].update(event["document_versions"])
        evidence["ids"].append(event["event_id"])
        if event["reason"].strip():
            evidence["reason"] = event["reason"]
    active = [row for row in rows if row["action_status"] != "none"]
    due = {row["canonical_job_id"]: local_date(row["action_due_at"]) for row in rows}
    counts = (
        ("total_applications", len(rows)),
        ("verified_jobs", sum(row["verification_status"] == "jd_verified" for row in rows)),
        ("approval_needs", sum(row["action_status"] == "candidate_review" for row in rows)),
        ("blocked_actions", sum(row["action_status"] == "blocked" for row in rows)),
        ("overdue_actions", sum(due[row["canonical_job_id"]] is not None
                               and due[row["canonical_job_id"]] < today for row in active)),
        ("due_today", sum(due[row["canonical_job_id"]] == today for row in active)),
    )
    section(overview, f"Operational dashboard - {today} - Australia/Melbourne",
            ("count", "value", "source_path"), [(name, count, SOURCES[0]) for name, count in counts])
    section(overview, f"Conversion summary - as of {summary['as_of']}",
            METRIC_HEADERS, metric_rows(summary["overall"]))
    priorities = []
    for row in active:
        date = due[row["canonical_job_id"]]
        categories = []
        if date is not None and date <= today:
            categories.append("due_today" if date == today else "overdue")
        if row["action_status"] in ("candidate_review", "blocked"):
            categories.append(row["action_status"])
        if categories:
            priorities.append((row["canonical_job_id"], date, row["next_action"],
                               ", ".join(categories), row["blocker_category"], SOURCES[0]))
    section(overview, "Priority actions", ("canonical_job_id", "action_due_at", "next_action",
                                         "category", "blocker_category", "source_path"), priorities)
    application_rows = []
    followup_rows = []
    for row in rows:
        identifier = row["canonical_job_id"]
        evidence = history[identifier]
        event_ids = ", ".join(evidence["ids"])
        values = [due[identifier] if field == "action_due_at" else row[field] for field in HEADER]
        application_rows.append((*values, json.dumps(evidence["versions"], ensure_ascii=False,
                                                      sort_keys=True), event_ids, *SOURCES[:2]))
        if row["action_status"] != "none":
            followup_rows.append((identifier, due[identifier], row["application_stage"],
                                  row["action_status"], row["next_action"], row["blocker_category"],
                                  evidence["reason"], event_ids, *SOURCES[:2]))
    section(applications, "Applications - current tracker state", (*HEADER, "document_versions",
            "event_ids", "tracker_reference", "event_reference"), application_rows, filtering=True)
    conversion_rows = [("overall", "overall", *values) for values in metric_rows(summary["overall"])]
    for source, segment in summary["segments"]["source"].items():
        conversion_rows.extend(("source", source, *values) for values in metric_rows(segment["metrics"]))
    section(conversion, f"Summary metrics - as of {summary['as_of']}",
            ("segment_dimension", "segment", *METRIC_HEADERS), conversion_rows, filtering=True)
    weeks = defaultdict(lambda: [0, 0, 0])
    for row in rows:
        for index, field in enumerate(("applied_at", "interview_at", "offer_at")):
            date = local_date(row[field])
            if date is not None:
                weeks[date - dt.timedelta(days=date.weekday())][index] += 1
    section(conversion, "Weekly counts - Monday starts - Australia/Melbourne",
            ("week_start", "applications", "interviews", "offers", "source_path"),
            [(week, *counts, SOURCES[0]) for week, counts in sorted(weeks.items())])
    section(followups, "Follow-ups - blank dates mean no due date recorded",
            ("canonical_job_id", "action_due_at", "application_stage", "action_status",
             "next_action", "blocker_category", "reason", "event_ids", "tracker_reference",
             "event_reference"), followup_rows, filtering=True)
    return workbook


def publish(root, destination, workbook):
    contained(root, destination)
    destination.parent.mkdir(exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w+b", dir=destination.parent,
                                         prefix=".dashboard-", suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
            workbook.save(handle)
            handle.flush()
            os.fsync(handle.fileno())
        contained(root, destination)
        contained(root, temporary)
        os.replace(temporary, destination)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", required=True, type=Path)
    arguments = parser.parse_args()
    try:
        root, paths = workspace_paths(arguments.workspace)
        jobs = read_rows(root, paths[0])
        events = read_events(root, paths[1], jobs)
        summary = read_summary(root, paths[2])
        workbook = build_workbook(jobs, events, summary)
        publish(root, paths[3], workbook)
    except (OSError, ValueError, csv.Error, OverflowError, RecursionError) as error:
        print(f"Dashboard build rejected: {error or type(error).__name__}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
